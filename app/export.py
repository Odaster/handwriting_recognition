"""Export recognized text to txt, Markdown, Word (docx), Excel (xlsx) or PDF."""

from __future__ import annotations

import io
import os

# (media_type, file extension) per format.
FORMATS = {
    "txt": ("text/plain; charset=utf-8", "txt"),
    "md": ("text/markdown; charset=utf-8", "md"),
    "docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
    "xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    ),
    "pdf": ("application/pdf", "pdf"),
}

_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\DejaVuSans.ttf",
    "/Library/Fonts/Arial.ttf",
]


def _paragraphs(text: str) -> list[str]:
    """Split on blank lines into paragraphs (preserving inner line breaks)."""
    blocks, current = [], []
    for line in (text or "").split("\n"):
        if line.strip():
            current.append(line)
        elif current:
            blocks.append("\n".join(current))
            current = []
    if current:
        blocks.append("\n".join(current))
    return blocks


def _find_font() -> str | None:
    for path in _FONT_CANDIDATES:
        if os.path.exists(path):
            return path
    return None


def _to_docx(text: str) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Экспорт в Word требует python-docx (см. requirements.txt).") from exc

    doc = Document()
    for block in _paragraphs(text):
        para = doc.add_paragraph()
        for i, line in enumerate(block.split("\n")):
            if i:
                para.add_run().add_break()
            para.add_run(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _to_xlsx(text: str) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("Экспорт в Excel требует openpyxl (см. requirements.txt).") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "OCR"
    ws.append(["#", "Текст"])
    row = 1
    for line in (text or "").split("\n"):
        if line.strip():
            ws.append([row, line])
            row += 1
    ws.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _to_pdf(text: str) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RuntimeError("Экспорт в PDF требует reportlab (см. requirements.txt).") from exc

    font_path = _find_font()
    font_name = "Helvetica"
    if font_path:
        font_name = "OCRFont"
        pdfmetrics.registerFont(TTFont(font_name, font_path))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    style = styles["Normal"]
    style.fontName = font_name
    style.fontSize = 12
    style.leading = 16

    story = []
    for block in _paragraphs(text):
        safe = block.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe.replace("\n", "<br/>"), style))
        story.append(Spacer(1, 10))
    if not story:
        story.append(Paragraph(" ", style))
    doc.build(story)
    return buf.getvalue()


def build_export(text: str, fmt: str) -> tuple[bytes, str, str]:
    """Return (bytes, media_type, extension) for the requested format."""
    if fmt not in FORMATS:
        raise ValueError(f"Unknown format: {fmt}")
    media_type, ext = FORMATS[fmt]

    if fmt == "txt":
        data = (text or "").encode("utf-8")
    elif fmt == "md":
        data = (text or "").encode("utf-8")
    elif fmt == "docx":
        data = _to_docx(text)
    elif fmt == "xlsx":
        data = _to_xlsx(text)
    elif fmt == "pdf":
        data = _to_pdf(text)
    else:  # pragma: no cover - guarded above
        raise ValueError(fmt)

    return data, media_type, ext
